import io

import pdfplumber
from docx import Document

from .resume_cleaner import normalize_resume_text


ALLOWED_EXTENSIONS = {
    ".pdf",
    ".docx",
}

ALLOWED_CONTENT_TYPES = {
    "application/pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


class ResumeParsingError(Exception):
    """Raised when resume parsing fails."""
    pass


def get_file_extension(filename):
    """
    Return the lowercase file extension.
    """

    if not filename:
        return ""

    filename = filename.lower()

    if "." not in filename:
        return ""

    return "." + filename.rsplit(".", 1)[1]


def validate_resume_file(file):
    """
    Validate uploaded resume file.
    """

    if not file:
        raise ResumeParsingError("Resume file is required.")

    extension = get_file_extension(file.name)

    if extension not in ALLOWED_EXTENSIONS:
        raise ResumeParsingError(
            "Unsupported file type. Only PDF and DOCX files are allowed."
        )

    # Content type validation is helpful, but some clients may send
    # application/octet-stream, so extension remains the primary check.
    if (
        file.content_type
        and file.content_type not in ALLOWED_CONTENT_TYPES
        and file.content_type != "application/octet-stream"
    ):
        raise ResumeParsingError(
            "Invalid file content type."
        )

    return extension


def extract_pdf_text(file):
    """
    Extract text from a PDF file using pdfplumber.
    """

    try:
        file.seek(0)

        with pdfplumber.open(file) as pdf:

            extracted_pages = []

            for page in pdf.pages:
                page_text = page.extract_text()

                if page_text:
                    extracted_pages.append(page_text)

            raw_text = "\n\n".join(extracted_pages)

    except Exception as exc:
        raise ResumeParsingError(
            f"Unable to extract text from PDF: {str(exc)}"
        )

    if not raw_text.strip():
        raise ResumeParsingError(
            "No readable text was found in the PDF."
        )

    return raw_text


def extract_docx_text(file):
    """
    Extract text from a DOCX file using python-docx.
    """

    try:
        file.seek(0)

        document = Document(file)

        paragraphs = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()

            if text:
                paragraphs.append(text)

        # Also extract text from tables because resumes often
        # contain skills, education, and experience in tables.
        for table in document.tables:

            for row in table.rows:

                row_values = []

                for cell in row.cells:
                    cell_text = cell.text.strip()

                    if cell_text:
                        row_values.append(cell_text)

                if row_values:
                    paragraphs.append(" | ".join(row_values))

        raw_text = "\n".join(paragraphs)

    except Exception as exc:
        raise ResumeParsingError(
            f"Unable to extract text from DOCX: {str(exc)}"
        )

    if not raw_text.strip():
        raise ResumeParsingError(
            "No readable text was found in the DOCX file."
        )

    return raw_text


def extract_resume_text(file):
    """
    Main resume extraction function.

    Detects PDF/DOCX and extracts text accordingly.
    """

    extension = validate_resume_file(file)

    if extension == ".pdf":
        raw_text = extract_pdf_text(file)

    elif extension == ".docx":
        raw_text = extract_docx_text(file)

    else:
        raise ResumeParsingError(
            "Unsupported resume format."
        )

    cleaned_text = normalize_resume_text(raw_text)

    if not cleaned_text:
        raise ResumeParsingError(
            "Resume text is empty after cleaning."
        )

    return {
        "file_type": extension.replace(".", "").upper(),
        "raw_text": raw_text,
        "cleaned_text": cleaned_text,
    }